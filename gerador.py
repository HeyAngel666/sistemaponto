"""Geração do cartão ponto em Excel.

As regras de cálculo (horas extras, faltas, atestados e feriados) são as
mesmas desde a primeira versão do sistema. O que mudou foi a origem dos
horários: agora vêm do cadastro da empresa (empresas.py).
"""

import calendar
import random
import re
import sys
from datetime import datetime, time, timedelta
from pathlib import Path

import holidays
from openpyxl import load_workbook

from empresas import (
    DOMINGO,
    NOME_DIA_SEMANA,
    cabecalho_empresa,
    espelho_de_jornada,
    horario_do_dia,
    obter_empresa,
)


# =====================================
# CONFIGURAÇÕES
# =====================================
def _pasta_do_programa():
    """Onde ficam os arquivos do sistema (dentro do .exe, quando empacotado)."""
    if getattr(sys, "frozen", False):
        return Path(sys._MEIPASS)
    return Path(__file__).resolve().parent


def _pasta_de_trabalho():
    """Onde os cartões são salvos (ao lado do .exe, quando empacotado)."""
    if getattr(sys, "frozen", False):
        return Path(sys.executable).resolve().parent
    return Path(__file__).resolve().parent


PASTA_MODELOS = _pasta_do_programa() / "modelo"
PASTA_SAIDA = _pasta_de_trabalho() / "cartoes"

# Onde a tabela de dias começa na planilha modelo
LINHA_INICIAL = 6
LINHA_FINAL = 36

# Células fixas do modelo
CELULA_EMPRESA = "B2"
CELULA_NOME = "B3"

# Coluna auxiliar do modelo: duração de cada hora extra, somada no total
COLUNA_DURACAO = "H"

# Colunas da tabela
COL_DATA = 1
COL_ENTRADA_MANHA = 2
COL_SAIDA_MANHA = 3
COL_ENTRADA_TARDE = 4
COL_SAIDA_TARDE = 5
COL_ENTRADA_EXTRA = 6
COL_SAIDA_EXTRA = 7
COL_LETRA_ENTRADA_EXTRA = "F"
COL_LETRA_SAIDA_EXTRA = "G"

# Limites de horas extras por dia (em minutos)
# 50%: lançadas após a jornada, nos dias úteis
EXTRA_MIN_POR_DIA = 30
EXTRA_MAX_POR_DIA = 120
# 100%: domingos e feriados trabalhados, até uma jornada inteira
EXTRA_100_MIN_POR_DIA = 30
EXTRA_100_MAX_POR_DIA = 480
# Abaixo disso não compensa parar para o almoço e voltar
MINIMO_PARA_VIRAR_TARDE = 60

FERIADOS_BR = holidays.Brazil()

MESES_PT = {
    1: "JANEIRO", 2: "FEVEREIRO", 3: "MARÇO", 4: "ABRIL",
    5: "MAIO", 6: "JUNHO", 7: "JULHO", 8: "AGOSTO",
    9: "SETEMBRO", 10: "OUTUBRO", 11: "NOVEMBRO", 12: "DEZEMBRO",
}

_TRADUCAO_FERIADOS = {
    "New Year's Day": "CONFRATERNIZAÇÃO UNIVERSAL",
    "New Year's Day (observed)": "CONFRATERNIZAÇÃO UNIVERSAL",
    "Carnival": "CARNAVAL",
    "Carnival Monday": "CARNAVAL",
    "Carnival Tuesday": "CARNAVAL",
    "Ash Wednesday": "QUARTA-FEIRA DE CINZAS",
    "Good Friday": "SEXTA-FEIRA SANTA",
    "Easter Sunday": "PÁSCOA",
    "Tiradentes' Day": "TIRADENTES",
    "Tiradentes Day": "TIRADENTES",
    "Worker's Day": "DIA DO TRABALHO",
    "Workers' Day": "DIA DO TRABALHO",
    "Labor Day": "DIA DO TRABALHO",
    "Corpus Christi": "CORPUS CHRISTI",
    "Independence Day": "INDEPENDÊNCIA DO BRASIL",
    "Our Lady of Aparecida": "NOSSA SENHORA APARECIDA",
    "All Souls' Day": "FINADOS",
    "All Souls Day": "FINADOS",
    "Republic Day": "PROCLAMAÇÃO DA REPÚBLICA",
    "Republic Proclamation Day": "PROCLAMAÇÃO DA REPÚBLICA",
    "Christmas Day": "NATAL",
    "Christmas": "NATAL",
}


def traduzir_feriado(nome):
    """Nome do feriado em português."""
    return _TRADUCAO_FERIADOS.get(nome, str(nome).upper())


# =====================================
# ESCRITA NA PLANILHA
# =====================================
def set_valor(ws, row, col, value):
    """Escreve respeitando células mescladas (escreve sempre na âncora)."""
    cell = ws.cell(row=row, column=col)
    for merged in ws.merged_cells.ranges:
        if cell.coordinate in merged:
            ws.cell(merged.min_row, merged.min_col).value = value
            return
    cell.value = value


def calcular_horario(hora_str, variar):
    """Converte 'HH:MM' em datetime, com variação opcional de ±5 minutos."""
    base = datetime.strptime(hora_str, "%H:%M")
    if variar:
        base += timedelta(minutes=random.randint(-5, 5))
    return base


# =====================================
# DISTRIBUIÇÃO DAS HORAS EXTRAS
# =====================================
def distribuir_exato(total_minutos, dias_validos,
                     minimo=EXTRA_MIN_POR_DIA, maximo=EXTRA_MAX_POR_DIA,
                     concentrar=False):
    """Distribui o total de minutos extras entre os dias disponíveis.

    Cada dia que recebe extra fica entre `minimo` e `maximo` minutos.

    `concentrar` usa o menor número de dias possível — é o caso do domingo
    trabalhado, que costuma ser um dia inteiro e não um pouco em cada.
    """
    dias_validos = list(dias_validos)

    if total_minutos <= 0:
        return {d: 0 for d in dias_validos}

    max_dias_possiveis = len(dias_validos)
    if max_dias_possiveis == 0:
        raise ValueError("Não há dias disponíveis para lançar horas extras.")

    if total_minutos > max_dias_possiveis * maximo:
        raise ValueError(
            "Horas extras excedem o máximo possível "
            f"({maximo} min/dia x {max_dias_possiveis} dias)."
        )

    min_dias_necessarios = -(-total_minutos // maximo)
    max_dias_necessarios = min(total_minutos // minimo, max_dias_possiveis)

    if max_dias_necessarios < min_dias_necessarios:
        raise ValueError(
            "Não é possível distribuir as horas extras respeitando o mínimo de "
            f"{minimo} min e o máximo de {maximo} min por dia."
        )

    qtd_dias = (
        int(min_dias_necessarios) if concentrar
        else random.randint(int(min_dias_necessarios), int(max_dias_necessarios))
    )
    dias_escolhidos = random.sample(dias_validos, qtd_dias)

    extras = {d: minimo for d in dias_escolhidos}
    restante = total_minutos - minimo * qtd_dias

    dias = dias_escolhidos[:]
    while restante > 0:
        progresso = False
        random.shuffle(dias)

        for d in dias:
            if restante <= 0:
                break

            livre = maximo - extras[d]
            if livre <= 0:
                continue

            valor = min(random.choice([5, 10, 15, 20, 30]), livre, restante)
            extras[d] += valor
            restante -= valor
            progresso = True

        if not progresso:
            raise ValueError("Erro ao distribuir as horas extras.")

    for d in dias_validos:
        extras.setdefault(d, 0)

    return extras


# =====================================
# GERAÇÃO DO CARTÃO
# =====================================
def montar_cartao(empresa_codigo, nome, mes, ano, horas_extras=0, faltas=0,
                  atestados=0, dia_inicio=1, dia_fim=None, horas_extras_100=0,
                  dias_ferias=0, inicio_ferias=None):
    """Monta o conteúdo do cartão (sem gravar arquivo).

    É a fonte única do que vai para o Excel e para o PDF.

    horas_extras      -> extras de 50%, lançadas após a jornada dos dias úteis
    horas_extras_100  -> extras de 100%, lançadas em domingos e feriados
    dias_ferias       -> dias corridos de férias, a partir de inicio_ferias

    dia_inicio / dia_fim delimitam o período ativo no mês (admissão e
    desligamento).
    """
    empresa = obter_empresa(empresa_codigo)
    validar_entrada(nome, mes, ano, horas_extras, faltas, atestados)

    if horas_extras_100 < 0:
        raise ValueError("Horas extras de 100% não podem ser negativas.")

    ultimo_dia_mes = calendar.monthrange(ano, mes)[1]
    dia_fim = ultimo_dia_mes if dia_fim in (None, "") else int(dia_fim)
    dia_inicio = int(dia_inicio or 1)

    if not 1 <= dia_inicio <= ultimo_dia_mes:
        raise ValueError(f"Dia inicial deve estar entre 1 e {ultimo_dia_mes}.")
    if not dia_inicio <= dia_fim <= ultimo_dia_mes:
        raise ValueError(
            f"Dia final deve estar entre {dia_inicio} e {ultimo_dia_mes}."
        )

    dias_de_ferias = _periodo_de_ferias(
        dias_ferias, inicio_ferias, dia_inicio, dia_fim, ultimo_dia_mes
    )

    dias_trabalhados = empresa["dias_trabalhados"]

    # Dias em que o funcionário deveria trabalhar (sem feriados e sem férias)
    dias_uteis = [
        d for d in range(dia_inicio, dia_fim + 1)
        if datetime(ano, mes, d).weekday() in dias_trabalhados
        and datetime(ano, mes, d).date() not in FERIADOS_BR
        and d not in dias_de_ferias
    ]

    if faltas + atestados > len(dias_uteis):
        raise ValueError(
            f"Faltas + atestados ({faltas + atestados}) excedem os dias úteis "
            f"do período ({len(dias_uteis)})."
        )

    disponiveis = dias_uteis[:]
    dias_falta = random.sample(disponiveis, min(faltas, len(disponiveis)))
    for d in dias_falta:
        disponiveis.remove(d)

    dias_atestado = random.sample(disponiveis, min(atestados, len(disponiveis)))

    dias_extras = [
        d for d in dias_uteis
        if d not in dias_falta and d not in dias_atestado
    ]
    extras = distribuir_exato(int(round(horas_extras * 60)), dias_extras)

    # Extras de 100%: domingos e feriados dentro do período (fora das férias)
    dias_100_possiveis = [
        d for d in range(dia_inicio, dia_fim + 1)
        if (datetime(ano, mes, d).weekday() == DOMINGO
            or datetime(ano, mes, d).date() in FERIADOS_BR)
        and d not in dias_de_ferias
    ]
    extras_100 = distribuir_exato(
        int(round(horas_extras_100 * 60)), dias_100_possiveis,
        minimo=EXTRA_100_MIN_POR_DIA, maximo=EXTRA_100_MAX_POR_DIA,
        concentrar=True,
    )

    admissao = datetime(ano, mes, dia_inicio).strftime("%d/%m")
    linhas = []

    for d in range(1, ultimo_dia_mes + 1):
        data = datetime(ano, mes, d)
        dia_semana = data.weekday()

        if d < dia_inicio:
            colunas = _ocorrencia(f"ADMITIDO EM {admissao}")
        elif d > dia_fim:
            colunas = _ocorrencia("DESLIGADO")
        elif d in dias_de_ferias:
            colunas = _ocorrencia("FÉRIAS")
        elif extras_100.get(d, 0) > 0:
            colunas = _marcacoes_de_cem_por_cento(data, empresa, extras_100[d])
        elif dia_semana == DOMINGO or dia_semana not in dias_trabalhados:
            colunas = _ocorrencia(NOME_DIA_SEMANA[dia_semana])
        elif data.date() in FERIADOS_BR:
            colunas = _ocorrencia(traduzir_feriado(FERIADOS_BR.get(data.date())))
        elif d in dias_falta:
            colunas = _ocorrencia("FALTA")
        elif d in dias_atestado:
            colunas = _ocorrencia("ATESTADO")
        else:
            colunas = _marcacoes_de_dia_trabalhado(data, empresa, extras.get(d, 0))

        linhas.append({"data": data.strftime("%d/%m"), "colunas": colunas})

    return {
        "empresa_codigo": str(empresa_codigo).strip().upper(),
        "empresa": cabecalho_empresa(empresa),
        "espelho": espelho_de_jornada(empresa) if empresa["espelho_de_jornada"] else None,
        "nome": nome,
        "competencia": f"{MESES_PT[mes]}/{ano}",
        "mes": mes,
        "ano": ano,
        "linhas": linhas,
    }


def _ocorrencia(texto):
    """Ocorrência ocupa as seis colunas de marcação."""
    return [texto] * 6


def _periodo_de_ferias(dias_ferias, inicio_ferias, dia_inicio, dia_fim, ultimo_dia_mes):
    dias_ferias = int(dias_ferias or 0)
    if dias_ferias <= 0:
        return set()

    if inicio_ferias in (None, "", 0):
        raise ValueError(
            f"São {dias_ferias} dia(s) de férias: informe o dia em que começam."
        )

    inicio_ferias = int(inicio_ferias)
    if not 1 <= inicio_ferias <= ultimo_dia_mes:
        raise ValueError(
            f"Início das férias deve estar entre 1 e {ultimo_dia_mes}."
        )

    return {
        d for d in range(inicio_ferias, inicio_ferias + dias_ferias)
        if dia_inicio <= d <= dia_fim
    }


def gerar_cartao(empresa_codigo, nome, mes, ano, horas_extras=0, faltas=0,
                 atestados=0, dia_inicio=1, dia_fim=None, pasta_saida=None,
                 horas_extras_100=0, dias_ferias=0, inicio_ferias=None):
    """Gera o cartão ponto em Excel e devolve o caminho do arquivo."""
    cartao = montar_cartao(
        empresa_codigo, nome, mes, ano, horas_extras, faltas, atestados,
        dia_inicio, dia_fim, horas_extras_100, dias_ferias, inicio_ferias,
    )
    return salvar_em_excel(cartao, pasta_saida)


def modelo_da_empresa(empresa_codigo):
    empresa = obter_empresa(empresa_codigo)
    caminho = PASTA_MODELOS / empresa["modelo"]
    if not caminho.exists():
        raise FileNotFoundError(f"Planilha modelo não encontrada: {caminho}")
    return caminho


def _como_hora(valor):
    """Converte 'HH:MM' em hora de verdade, para as fórmulas do modelo somarem."""
    if isinstance(valor, str) and re.fullmatch(r"\d{1,2}:\d{2}", valor):
        horas, minutos = valor.split(":")
        return time(int(horas), int(minutos))
    return valor


def salvar_em_excel(cartao, pasta_saida=None):
    """Grava em Excel um cartão já montado.

    Separado da montagem para que o Excel e o PDF do mesmo funcionário
    saiam com as mesmas marcações.
    """
    mes, ano, nome = cartao["mes"], cartao["ano"], cartao["nome"]

    wb = load_workbook(modelo_da_empresa(cartao["empresa_codigo"]))
    ws = wb.active

    ws[CELULA_EMPRESA] = cartao["empresa"]
    ws[CELULA_NOME] = cartao["nome"]

    for indice, registro in enumerate(cartao["linhas"]):
        linha = LINHA_INICIAL + indice
        set_valor(ws, linha, COL_DATA, registro["data"])
        for coluna, valor in enumerate(registro["colunas"], start=COL_ENTRADA_MANHA):
            set_valor(ws, linha, coluna, _como_hora(valor))

    # Nos dias de ocorrência as colunas de extra levam texto (DOMINGO, FALTA...),
    # e a subtração da coluna auxiliar daria erro: o IFERROR mantém o total certo.
    for linha in range(LINHA_INICIAL, LINHA_FINAL + 1):
        ws[f"{COLUNA_DURACAO}{linha}"] = (
            f"=IFERROR({COL_LETRA_SAIDA_EXTRA}{linha}-"
            f"{COL_LETRA_ENTRADA_EXTRA}{linha},0)"
        )

    # Sem pasta escolhida, organiza por mês dentro da pasta padrão
    pasta = Path(pasta_saida) if pasta_saida else PASTA_SAIDA / f"{ano}-{mes:02d}"
    pasta.mkdir(parents=True, exist_ok=True)

    caminho = pasta / f"{nome_arquivo_seguro(nome)} - {mes:02d}-{ano}.xlsx"
    wb.save(caminho)
    return caminho


def _marcacoes_de_dia_trabalhado(data, empresa, extra_min):
    """As seis marcações de um dia normal de trabalho."""
    horario = horario_do_dia(empresa, data.weekday())
    variacoes = empresa["variacoes"]

    marcacoes = {
        campo: calcular_horario(horario[campo], variacoes.get(campo, False))
        for campo in ("entrada", "saida_almoco", "volta_almoco", "saida")
    }

    colunas = [
        marcacoes["entrada"].strftime("%H:%M"),
        marcacoes["saida_almoco"].strftime("%H:%M"),
        marcacoes["volta_almoco"].strftime("%H:%M"),
        marcacoes["saida"].strftime("%H:%M"),
    ]

    if extra_min > 0:
        fim_extra = marcacoes["saida"] + timedelta(minutes=extra_min)
        colunas += [marcacoes["saida"].strftime("%H:%M"), fim_extra.strftime("%H:%M")]
    else:
        colunas += ["", ""]

    return colunas


def _marcacoes_de_cem_por_cento(data, empresa, minutos):
    """Domingo ou feriado trabalhado: marcação normal, hora extra de 100%.

    Começa no horário de entrada da empresa e segue pelas horas lançadas,
    passando para o período da tarde quando ultrapassa o horário de almoço.
    """
    horario = horario_do_dia(empresa, data.weekday())
    variacoes = empresa["variacoes"]

    entrada = calcular_horario(horario["entrada"], variacoes.get("entrada", False))
    saida_almoco = datetime.strptime(horario["saida_almoco"], "%H:%M")
    volta_almoco = datetime.strptime(horario["volta_almoco"], "%H:%M")

    minutos_da_manha = int((saida_almoco - entrada).total_seconds() // 60)

    # Sobra pequena não vale parar para o almoço: emenda tudo no período da manhã
    if minutos <= minutos_da_manha + MINIMO_PARA_VIRAR_TARDE:
        fim = entrada + timedelta(minutes=minutos)
        return [entrada.strftime("%H:%M"), fim.strftime("%H:%M"), "", "", "", ""]

    fim = volta_almoco + timedelta(minutes=minutos - minutos_da_manha)
    return [
        entrada.strftime("%H:%M"),
        saida_almoco.strftime("%H:%M"),
        volta_almoco.strftime("%H:%M"),
        fim.strftime("%H:%M"),
        "",
        "",
    ]


def validar_entrada(nome, mes, ano, horas_extras, faltas, atestados):
    if not str(nome).strip():
        raise ValueError("Informe o nome do funcionário.")
    if mes not in MESES_PT:
        raise ValueError("Mês deve ser um número de 1 a 12.")
    if not 2000 <= ano <= 2100:
        raise ValueError("Ano inválido.")
    if horas_extras < 0:
        raise ValueError("Horas extras não podem ser negativas.")
    if faltas < 0 or atestados < 0:
        raise ValueError("Faltas e atestados não podem ser negativos.")


def nome_arquivo_seguro(nome):
    """Remove caracteres que o Windows não aceita em nome de arquivo."""
    invalidos = '<>:"/\\|?*'
    limpo = "".join("-" if c in invalidos else c for c in str(nome))
    return " ".join(limpo.split()).strip(". ")
