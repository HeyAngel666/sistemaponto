"""Geração em lote a partir de uma planilha de funcionários.

Fluxo: o usuário gera a planilha em branco, preenche o quadro de
funcionários do mês e importa de volta para gerar todos os cartões.
"""

import calendar
import unicodedata
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

from empresas import listar_empresas

COLUNAS = [
    ("EMPRESA", "MONTSUL ou VAGNER"),
    ("NOME", "Nome do funcionário"),
    ("MES", "1 a 12"),
    ("ANO", "Ex: 2026"),
    ("HORAS EXTRAS", "Extras de 50% (dias úteis). Deixe 0 se não houver"),
    ("HORAS EXTRAS 100", "Extras de 100% (domingos/feriados). Deixe 0 se não houver"),
    ("FALTAS", "Quantidade de faltas no mês"),
    ("ATESTADOS", "Quantidade de atestados no mês"),
    ("DIA INICIAL", "Deixe vazio se trabalhou o mês todo"),
    ("DIA FINAL", "Preencha só se foi desligado no meio do mês"),
    ("FERIAS DIAS", "Dias corridos de férias no mês. Deixe 0 se não houver"),
    ("FERIAS INICIO", "Dia em que as férias começam"),
]


def normalizar(texto):
    """Remove acentos, espaços extras e deixa em maiúsculo."""
    return (
        unicodedata.normalize("NFD", str(texto))
        .encode("ascii", "ignore")
        .decode("ascii")
        .upper()
        .strip()
    )


def criar_planilha_modelo(caminho, registros=None):
    """Cria a planilha de funcionários.

    Sem registros, sai em branco (com uma linha de exemplo) para o usuário
    preencher. Com registros — vindos do holerite — já sai preenchida.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Funcionários"

    fonte = "Arial"
    thin = Side(style="thin", color="9E9E9E")
    borda = Border(left=thin, right=thin, top=thin, bottom=thin)

    titulo = ws.cell(row=1, column=1, value="PREENCHA UMA LINHA POR FUNCIONÁRIO")
    titulo.font = Font(name=fonte, size=12, bold=True)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(COLUNAS))

    ajuda = ws.cell(
        row=2, column=1,
        value="Confira os dados e ajuste o que for preciso antes de importar."
        if registros else
        "A linha em amarelo é só um exemplo — apague antes de importar.",
    )
    ajuda.font = Font(name=fonte, size=10, italic=True, color="757575")
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(COLUNAS))

    linha_cabecalho = 4
    for idx, (nome_coluna, descricao) in enumerate(COLUNAS, start=1):
        cel = ws.cell(row=linha_cabecalho, column=idx, value=nome_coluna)
        cel.font = Font(name=fonte, size=10, bold=True, color="FFFFFF")
        cel.fill = PatternFill("solid", fgColor="37474F")
        cel.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cel.border = borda

        dica = ws.cell(row=linha_cabecalho + 1, column=idx, value=descricao)
        dica.font = Font(name=fonte, size=9, italic=True, color="757575")
        dica.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        dica.border = borda

        ws.column_dimensions[cel.column_letter].width = max(14, len(nome_coluna) + 4)

    primeira_linha_livre = linha_cabecalho + 2

    if registros:
        for registro in registros:
            valores = [
                registro.get("empresa_codigo") or "",
                registro.get("nome", ""),
                registro.get("mes"),
                registro.get("ano"),
                registro.get("horas_extras", 0),
                registro.get("horas_extras_100", 0),
                registro.get("faltas", 0),
                registro.get("atestados", 0),
                registro.get("dia_inicio") if registro.get("dia_inicio", 1) != 1 else None,
                registro.get("dia_fim"),
                registro.get("dias_ferias", 0),
                registro.get("inicio_ferias"),
            ]
            for idx, valor in enumerate(valores, start=1):
                cel = ws.cell(row=primeira_linha_livre, column=idx, value=valor)
                cel.font = Font(name=fonte, size=10)
                cel.alignment = Alignment(horizontal="center", vertical="center")
                cel.border = borda
            primeira_linha_livre += 1
    else:
        exemplo = ["MONTSUL", "JOÃO DA SILVA", 8, 2026, 12, 0, 1, 0, None, None, 0, None]
        for idx, valor in enumerate(exemplo, start=1):
            cel = ws.cell(row=primeira_linha_livre, column=idx, value=valor)
            cel.font = Font(name=fonte, size=10)
            cel.fill = PatternFill("solid", fgColor="FFF9C4")
            cel.alignment = Alignment(horizontal="center", vertical="center")
            cel.border = borda
        primeira_linha_livre += 1

    for row in range(primeira_linha_livre, primeira_linha_livre + 40):
        for col in range(1, len(COLUNAS) + 1):
            cel = ws.cell(row=row, column=col)
            cel.font = Font(name=fonte, size=10)
            cel.alignment = Alignment(horizontal="center", vertical="center")
            cel.border = borda

    empresas = ",".join(listar_empresas())
    validacao = DataValidation(type="list", formula1=f'"{empresas}"', allow_blank=True)
    ws.add_data_validation(validacao)
    validacao.add(f"A{linha_cabecalho + 2}:A{primeira_linha_livre + 40}")

    ws.row_dimensions[linha_cabecalho].height = 24
    ws.row_dimensions[linha_cabecalho + 1].height = 30
    ws.freeze_panes = ws.cell(row=linha_cabecalho + 2, column=1)

    caminho = Path(caminho)
    wb.save(caminho)
    return caminho


def ler_planilha(caminho):
    """Lê a planilha preenchida e devolve a lista de funcionários.

    Ignora a linha de exemplo e linhas em branco.
    """
    df = pd.read_excel(caminho, skiprows=3)
    df.columns = [normalizar(c) for c in df.columns]

    obrigatorias = {"EMPRESA", "NOME", "MES", "ANO"}
    faltando = obrigatorias - set(df.columns)
    if faltando:
        raise ValueError(
            "A planilha não tem as colunas: " + ", ".join(sorted(faltando))
        )

    # Descarta a linha de dica e a linha de exemplo
    df = df[df["NOME"].notna()]
    df = df[~df["NOME"].astype(str).str.upper().str.strip().isin(
        ["JOÃO DA SILVA", "NOME DO FUNCIONÁRIO"]
    )]

    registros = []
    for numero_linha, row in enumerate(df.to_dict("records"), start=1):
        try:
            mes = inteiro(row.get("MES"), obrigatorio=True)
            ano = inteiro(row.get("ANO"), obrigatorio=True)
            ultimo_dia = calendar.monthrange(ano, mes)[1]

            registros.append({
                "empresa_codigo": normalizar(row.get("EMPRESA")),
                "nome": str(row.get("NOME")).strip(),
                "mes": mes,
                "ano": ano,
                "horas_extras": decimal(row.get("HORAS EXTRAS"), padrao=0),
                "horas_extras_100": decimal(row.get("HORAS EXTRAS 100"), padrao=0),
                "faltas": inteiro(row.get("FALTAS"), padrao=0),
                "atestados": inteiro(row.get("ATESTADOS"), padrao=0),
                "dia_inicio": inteiro(row.get("DIA INICIAL"), padrao=1),
                "dia_fim": inteiro(row.get("DIA FINAL"), padrao=ultimo_dia),
                "dias_ferias": inteiro(row.get("FERIAS DIAS"), padrao=0),
                "inicio_ferias": inteiro(row.get("FERIAS INICIO"), padrao=None),
            })
        except (ValueError, TypeError) as erro:
            raise ValueError(f"Linha {numero_linha} da planilha: {erro}") from erro

    if not registros:
        raise ValueError("Nenhum funcionário preenchido na planilha.")

    return registros


def vazio(valor):
    return valor is None or pd.isna(valor) or str(valor).strip() == ""


def inteiro(valor, padrao=None, obrigatorio=False):
    if vazio(valor):
        if obrigatorio:
            raise ValueError("campo obrigatório em branco")
        return padrao
    return int(float(valor))


def decimal(valor, padrao=0.0):
    if vazio(valor):
        return padrao
    return float(str(valor).replace(",", "."))
